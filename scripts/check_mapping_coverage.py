# scripts/check_mapping_coverage.py
# Dry-run: apply product_name_mapping.json to the real workbook and report how much
# of the actual data maps to an internal SKU — BEFORE generating any SQL.
# Jalankan: python scripts/check_mapping_coverage.py Stock_Opname.xlsx
import sys
import json
import re
from collections import defaultdict
from openpyxl import load_workbook

mapping = json.load(open("scripts/product_name_mapping.json"))
name_aliases = mapping["name_aliases"]
code_aliases = mapping.get("code_aliases", {})


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


name_norm = {norm(k): v for k, v in name_aliases.items()}
code_norm = {norm(k): v for k, v in code_aliases.items()}


def sku_by_name(name):
    return name_norm.get(norm(name)) if name is not None else None


def sku_by_code(code):
    return code_norm.get(norm(code)) if code is not None else None


def col(r, i):
    return r[i] if i < len(r) else None


IN_SHEETS = ["Masuk Apr26", "Masuk May26", "Masuk June26"]
OUT_SHEETS = ["Keluar Apr26", "Keluar May26", "Keluar June26"]

wb = load_workbook(sys.argv[1] if len(sys.argv) > 1 else "Stock_Opname.xlsx",
                   read_only=True, data_only=True)

# --- transaction coverage ---
rows = 0
mapped = 0
unmapped = defaultdict(lambda: [0, 0])
txn_skus = set()
for sheet in IN_SHEETS + OUT_SHEETS:
    ws = wb[sheet]
    for r in ws.iter_rows(values_only=True):
        if col(r, 0) is not None and hasattr(col(r, 0), "strftime") and col(r, 1) is not None and col(r, 2) is not None:
            rows += 1
            name = str(col(r, 1)).strip()
            s = sku_by_name(name)
            if s:
                mapped += 1
                txn_skus.add(s)
            else:
                unmapped[name][0] += 1
                unmapped[name][1] += int(col(r, 2))

print(f"TRANSAKSI: total={rows} mapped={mapped} unmapped={rows - mapped} "
      f"coverage={mapped / rows * 100:.1f}%")
if unmapped:
    print("  -- transaksi TAK terpetakan (tambahkan ke name_aliases):")
    for name, (c, q) in sorted(unmapped.items(), key=lambda x: -x[1][0]):
        print(f"     [{c}x qty={q}] {name!r}")

# --- price-list coverage ---
ws = wb["coret coret"]
priced = []
for r in ws.iter_rows(values_only=True):
    if col(r, 0) is not None and isinstance(col(r, 0), (int, float)):
        code, name = col(r, 1), col(r, 2)
        cost, sell = col(r, 4), col(r, 6)
        sku = sku_by_code(code) or sku_by_name(name)
        priced.append((code, name, cost, sell, sku))

matched = sum(1 for p in priced if p[4])
print(f"\nPRICE LIST: rows={len(priced)} matched_to_SKU={matched} unmatched={len(priced) - matched}")
for code, name, cost, sell, sku in priced:
    if not sku:
        print(f"  NO SKU: code={code!r} name={name!r}")

# --- SKUs that have transactions but no price ---
priced_skus = {p[4] for p in priced if p[4]}
no_price = txn_skus - priced_skus
print(f"\nSKU dengan transaksi: {len(txn_skus)} | punya harga: {len(txn_skus & priced_skus)} | TANPA harga: {len(no_price)}")
for s in sorted(no_price):
    print(f"  tanpa-harga: {s}")
