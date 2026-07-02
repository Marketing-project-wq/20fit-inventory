# scripts/parse_excel.py
# Jalankan dengan: python scripts/parse_excel.py Stock_Opname.xlsx
import sys
from openpyxl import load_workbook
from datetime import datetime


def safe_date(v):
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    return str(v) if v else None


def col(r, i):
    """Safe cell access — real sheets have ragged (variable-length) rows."""
    return r[i] if i < len(r) else None


def parse_stock_opname(filepath):
    wb = load_workbook(filepath, read_only=True, data_only=True)

    # === 1. PRODUK dari sheet "coret coret" (ada harga aktual) ===
    products_with_price = []
    ws = wb['coret coret']
    for r in ws.iter_rows(values_only=True):
        if r[0] is not None and isinstance(r[0], (int, float)):
            # Kolom: NO | PRODUCT_CODE | NAMA | USD_CENTR | IDR_CENTR | USD_20FIT | IDR_20FIT
            products_with_price.append({
                'no': r[0],
                'product_code': r[1],
                'name': r[2],
                'cost_price_idr': float(r[4]) if r[4] else None,    # harga modal (harga CENTR)
                'selling_price_idr': float(r[6]) if r[6] else None,  # harga jual (harga 20FIT)
            })
    print(f"Produk dengan harga: {len(products_with_price)}")

    # === 2. SEMUA PRODUK dari LIST BARANG ===
    all_products = []
    ws = wb['LIST BARANG']
    for r in ws.iter_rows(values_only=True):
        if r[1] is not None and r[1] != 'Nama Barang':
            all_products.append(str(r[1]).strip())
    print(f"Total produk di LIST BARANG: {len(all_products)}")

    # === 3. TRANSAKSI MASUK ===
    goods_in = []
    for sheet_name in ['Masuk Apr26', 'Masuk May26', 'Masuk June26']:
        ws = wb[sheet_name]
        for r in ws.iter_rows(values_only=True):
            if col(r, 0) is not None and hasattr(col(r, 0), 'strftime') and col(r, 1) is not None and col(r, 2) is not None:
                goods_in.append({
                    'date': safe_date(col(r, 0)),
                    'product_name': str(col(r, 1)).strip(),
                    'quantity': int(col(r, 2)),
                    'notes': str(col(r, 3)).strip() if col(r, 3) else None,
                    'source_sheet': sheet_name,
                })
    print(f"Total transaksi masuk: {len(goods_in)}")

    # === 4. TRANSAKSI KELUAR ===
    goods_out = []
    for sheet_name in ['Keluar Apr26', 'Keluar May26', 'Keluar June26']:
        ws = wb[sheet_name]
        for r in ws.iter_rows(values_only=True):
            if col(r, 0) is not None and hasattr(col(r, 0), 'strftime') and col(r, 1) is not None and col(r, 2) is not None:
                # Kolom: Tanggal | Barang | Jumlah Keluar | Catatan/Sales | Nama Client
                goods_out.append({
                    'date': safe_date(col(r, 0)),
                    'product_name': str(col(r, 1)).strip(),
                    'quantity': int(col(r, 2)),
                    'salesperson': str(col(r, 3)).strip() if col(r, 3) else None,
                    'client': str(col(r, 4)).strip() if col(r, 4) else None,
                    'source_sheet': sheet_name,
                })
    print(f"Total transaksi keluar: {len(goods_out)}")

    return {
        'products_with_price': products_with_price,
        'all_products': all_products,
        'goods_in': goods_in,
        'goods_out': goods_out,
    }


if __name__ == '__main__':
    filepath = sys.argv[1] if len(sys.argv) > 1 else 'Stock_Opname.xlsx'
    data = parse_stock_opname(filepath)

    # Print summary
    print("\n=== RINGKASAN ===")
    print(f"Produk dengan harga: {len(data['products_with_price'])}")
    print(f"Total produk di LIST BARANG: {len(data['all_products'])}")
    print(f"Transaksi masuk: {len(data['goods_in'])}")
    print(f"Transaksi keluar: {len(data['goods_out'])}")

    # Cari produk yang ada di transaksi tapi tidak ada di price list
    priced_names = {p['name'].lower().strip() for p in data['products_with_price'] if p['name']}
    all_transacted = set()
    for t in data['goods_in'] + data['goods_out']:
        all_transacted.add(t['product_name'].lower())

    print("\n=== PRODUK DALAM TRANSAKSI TAPI BELUM ADA HARGA ===")
    for name in sorted(all_transacted - priced_names):
        print(f"  - {name}")
